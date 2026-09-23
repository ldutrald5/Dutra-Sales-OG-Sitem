import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def category_from_heading(heading):
    value = normalize(heading)
    mapping = {
        "roi": "Diagnóstico e ROI",
        "fluxo de conversa": "Abordagens",
        "personas": "Segmentos e personas",
        "mensagens": "Abordagens",
        "regras da call": "Regras da Call AI",
        "validado": "Governança",
        "crm": "Follow-up e CRM",
        "objec": "Objeções",
        "pos venda": "Pós-venda",
    }
    return next((category for key, category in mapping.items() if key in value), "Conhecimento comercial")


def import_jsonl(path):
    records, errors = [], []
    with path.open("r", encoding="utf-8-sig") as handle:
        for line_number, raw in enumerate(handle, 1):
            if not raw.strip():
                continue
            try:
                row = json.loads(raw)
                if not isinstance(row, dict):
                    raise ValueError("esperado objeto JSON")
                missing = [key for key in ("id", "category", "title", "statement", "source") if not row.get(key)]
                if missing:
                    raise ValueError("campos ausentes: " + ", ".join(missing))
                records.append({
                    "id": str(row["id"]),
                    "category": str(row["category"]),
                    "title": str(row["title"]),
                    "text": str(row["statement"]),
                    "status": str(row.get("status") or "Não informado"),
                    "confidence": str(row.get("confidence") or "Não informada"),
                    "tags": [str(tag) for tag in row.get("tags", []) if str(tag).strip()],
                    "source": str(row["source"]),
                    "origin": path.name,
                    "locator": f"linha {line_number}",
                })
            except Exception as exc:
                errors.append({"line": line_number, "error": str(exc)})
    return records, errors


def import_docx(path, canonical_texts):
    document = Document(path)
    records = []
    heading = "Introdução"
    counter = 0
    for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").casefold()
        if style.startswith(("heading", "título")):
            heading = text
            continue
        if style == "title" or normalize(text) in {"og sales brain", "base de conhecimento comercial v0 1"}:
            continue
        if normalize(text) in canonical_texts or normalize(text).startswith("fonte "):
            continue
        counter += 1
        records.append({
            "id": f"OG-DOCX-P{paragraph_number:03d}",
            "category": category_from_heading(heading),
            "title": heading,
            "text": text,
            "status": "Documento interno",
            "confidence": "Consultar fonte",
            "tags": normalize(heading).split()[:8],
            "source": path.name,
            "origin": path.name,
            "locator": f"parágrafo {paragraph_number}",
        })
    for table_number, table in enumerate(document.tables, 1):
        for row_number, row in enumerate(table.rows, 1):
            cells = [cell.text.strip() for cell in row.cells]
            text = " | ".join(value for value in cells if value)
            if not text or normalize(text) in canonical_texts:
                continue
            counter += 1
            records.append({
                "id": f"OG-DOCX-T{table_number:02d}R{row_number:02d}",
                "category": "Tabela de referência",
                "title": f"Tabela {table_number}",
                "text": text,
                "status": "Documento interno",
                "confidence": "Consultar fonte",
                "tags": ["tabela", "referência"],
                "source": path.name,
                "origin": path.name,
                "locator": f"tabela {table_number}, linha {row_number}",
            })
    return records


def main():
    jsonl, docx, output = map(Path, sys.argv[1:4])
    canonical, errors = import_jsonl(jsonl)
    canonical_texts = {normalize(record["text"]) for record in canonical}
    supplemental = import_docx(docx, canonical_texts)
    all_records = canonical + supplemental
    seen, duplicates = {}, []
    for record in all_records:
        fingerprint = normalize(record["text"])
        if fingerprint in seen:
            duplicates.append({"kept": seen[fingerprint], "duplicate": record["id"]})
        else:
            seen[fingerprint] = record["id"]
    deduplicated = [record for record in all_records if seen.get(normalize(record["text"])) == record["id"]]
    cautions = [
        record["id"] for record in deduplicated
        if any(term in normalize(record.get("status")) for term in ("requer validacao", "premissa"))
        or any(term in normalize(record["text"]) for term in ("nao garantir", "nao prometer"))
    ]
    result = {
        "schemaVersion": 1,
        "version": "0.1",
        "importedAt": datetime.now(timezone.utc).isoformat(),
        "sources": [
            {"file": jsonl.name, "sha256": sha256(jsonl)},
            {"file": docx.name, "sha256": sha256(docx)},
        ],
        "stats": {
            "jsonlRecords": len(canonical),
            "docxRecords": len(supplemental),
            "indexedRecords": len(deduplicated),
            "errors": len(errors),
            "duplicates": len(duplicates),
            "cautions": len(cautions),
        },
        "errors": errors,
        "duplicates": duplicates,
        "cautions": cautions,
        "records": deduplicated,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: result[key] for key in ("version", "stats", "sources")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

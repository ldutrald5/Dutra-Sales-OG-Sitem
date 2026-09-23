import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document


def normalized(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def inspect_jsonl(path):
    rows, errors = [], []
    with path.open("r", encoding="utf-8-sig") as handle:
        for number, raw in enumerate(handle, 1):
            if not raw.strip():
                continue
            try:
                value = json.loads(raw)
                if not isinstance(value, dict):
                    raise ValueError("a linha não contém um objeto JSON")
                rows.append((number, value))
            except Exception as exc:
                errors.append({"line": number, "error": str(exc)})
    keys = Counter(key for _, row in rows for key in row)
    text_candidates = []
    for number, row in rows:
        for key in ("text", "content", "texto", "body", "answer", "resposta"):
            if row.get(key):
                text_candidates.append((number, key, normalized(row[key])))
                break
    duplicates = [
        {"lines": [n for n, _, candidate in text_candidates if candidate == text], "text": text[:160]}
        for text, count in Counter(item[2] for item in text_candidates).items()
        if text and count > 1
    ]
    return rows, {"rows": len(rows), "errors": errors, "keys": dict(keys), "duplicates": duplicates}


def inspect_docx(path):
    document = Document(path)
    blocks = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            blocks.append({"kind": "paragraph", "index": index, "style": paragraph.style.name, "text": text})
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append({"kind": "table", "index": f"{table_index}:{row_index}", "style": "table", "text": " | ".join(cells)})
    seen = Counter(normalized(block["text"]) for block in blocks)
    duplicates = [text[:160] for text, count in seen.items() if text and count > 1]
    return blocks, {"blocks": len(blocks), "tables": len(document.tables), "duplicates": duplicates}


def main():
    jsonl = Path(sys.argv[1])
    docx = Path(sys.argv[2])
    rows, json_report = inspect_jsonl(jsonl)
    blocks, docx_report = inspect_docx(docx)
    report = {
        "jsonl": json_report,
        "docx": docx_report,
        "sha256": {
            jsonl.name: hashlib.sha256(jsonl.read_bytes()).hexdigest(),
            docx.name: hashlib.sha256(docx.read_bytes()).hexdigest(),
        },
        "jsonl_samples": [{"line": number, "value": value} for number, value in rows[:4]],
        "docx_outline": [block for block in blocks if block["style"].lower().startswith(("title", "heading", "título"))][:40],
        "docx_samples": blocks[:20],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
